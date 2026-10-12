import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Paleta de Cores
    BG_DARK = RGBColor(15, 23, 42)      # #0F172A Slate 900
    CARD_BG = RGBColor(30, 41, 59)      # #1E293B Slate 800
    CARD_BORDER = RGBColor(51, 65, 85)  # #334155 Slate 700
    ACCENT_BLUE = RGBColor(56, 189, 248) # #38BDF8 Sky 400
    ACCENT_CYAN = RGBColor(6, 182, 212) # #06B6D4 Cyan 500
    ACCENT_EMERALD = RGBColor(16, 185, 129) # #10B981 Emerald 500
    ACCENT_GOLD = RGBColor(245, 158, 11) # #F59E0B Amber 500
    TEXT_LIGHT = RGBColor(248, 250, 252) # #F8FAFC Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8 Slate 400

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_DARK

    def add_header(slide, tag_text, title_text):
        # Tag Badge
        tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
        tf = tag.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = tag_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

        # Title
        title = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf2 = title.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_LIGHT

    def add_card(slide, left, top, width, height, title, content_list, accent_color=ACCENT_BLUE):
        # Card Background Shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        shape.line.width = Pt(1)

        # Content Box
        txBox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = txBox.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = accent_color
        p.space_after = Pt(10)

        for item in content_list:
            p2 = tf.add_paragraph()
            p2.text = "• " + item
            p2.font.size = Pt(12)
            p2.font.color.rgb = TEXT_MUTED
            p2.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 1: CAPA (TITLE SLIDE)
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide1)

    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.1))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = CARD_BORDER
    card1.line.width = Pt(1.5)

    tb1 = slide1.shapes.add_textbox(Inches(1.5), Inches(1.6), Inches(10.333), Inches(4.3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "INSURMINDS — PROJETO FINAL (I2A2)  |  GRUPO IA JOURNEY"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "D&O Intelligence"
    p2.font.size = Pt(40)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_LIGHT
    p2.space_after = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "Plataforma Inteligente para Análise, Extração RAG e Comparação de Apólices D&O"
    p3.font.size = Pt(18)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_after = Pt(30)

    p4 = tf1.add_paragraph()
    p4.text = "Grupo: IA Journey  |  Autor: Arsénio António Monjane  |  monjearse@hotmail.com  |  Outubro 2026"
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = ACCENT_CYAN

    # -------------------------------------------------------------
    # SLIDE 2: O PROBLEMA
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide2)
    add_header(slide2, "Contexto de Mercado & Desafio", "O Problema da Análise Tradicional de Apólices D&O")

    add_card(slide2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "1. Alta Complexidade", 
             [
                 "Contratos extensos de 50 a 200+ páginas.",
                 "Linguagem jurídica densa e especializada.",
                 "Múltiplas coberturas (Side A, B, C), sublimites e franquias.",
                 "Dificuldade de localizar cláusulas críticas rapidamente."
             ], ACCENT_GOLD)

    add_card(slide2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "2. Ineficiência Operacional", 
             [
                 "3 a 6 horas de análise por apólice por especialista.",
                 "Custo operacional elevado para corretoras e seguradoras.",
                 "Gargalo no processo de underwriting e cotações.",
                 "Baixa capacidade de análise comparativa em lote."
             ], ACCENT_GOLD)

    add_card(slide2, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "3. Risco & Assimetria", 
             [
                 "Risco de omissão de sublimites ou franquias severas.",
                 "Cláusulas de exclusão (Atos Dolosos, Poluição, Severabilidade).",
                 "Falta de padronização nas matrizes comparativas.",
                 "Ausência de rastreabilidade direta no documento original."
             ], ACCENT_GOLD)

    # -------------------------------------------------------------
    # SLIDE 3: A SOLUÇÃO
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide3)
    add_header(slide3, "Proposta de Valor", "D&O Intelligence — A Solução Baseada em IA Agêntica")

    add_card(slide3, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "Automação Inteligente", 
             [
                 "Ingestão e parsing de PDF/Imagem em segundos.",
                 "OCR avançado PyMuPDF com preservação de layout.",
                 "Extração estruturada de campos via Pydantic Agent.",
                 "Redução do tempo de análise em mais de 95%."
             ], ACCENT_EMERALD)

    add_card(slide3, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "100% Grounded Proof", 
             [
                 "Rastreabilidade auditável de cada dado extraído.",
                 "Vinculação direta à página exata do PDF original.",
                 "Visualizador split-screen interativo na UI.",
                 "Eliminação do risco de alucinações de IA."
             ], ACCENT_EMERALD)

    add_card(slide3, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "Benchmark & Severidade", 
             [
                 "Comparação semântica automatizada (Apólice A vs B).",
                 "Matriz Delta de Riscos com pontuação de severidade.",
                 "Destaque de desvios em Limites, Franquias e Cláusulas.",
                 "Relatórios executivos prontos para tomada de decisão."
             ], ACCENT_EMERALD)

    # -------------------------------------------------------------
    # SLIDE 4: ARQUITETURA TÉCNICA E FERRAMENTAS
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide4)
    add_header(slide4, "Arquitetura do Sistema & Stack", "Arquitetura Modular Desacoplada e Ferramentas de IA")

    add_card(slide4, Inches(0.8), Inches(1.8), Inches(2.7), Inches(5.0), 
             "Frontend SPA", 
             [
                 "React 18 & Vite 5",
                 "TailwindCSS B2B Design System",
                 "Dark Mode Nativo",
                 "Split-Screen PDF Viewer",
                 "Recharts (Gráficos Radar)"
             ], ACCENT_BLUE)

    add_card(slide4, Inches(3.8), Inches(1.8), Inches(2.7), Inches(5.0), 
             "Backend REST", 
             [
                 "FastAPI 0.115 & Python 3.12+",
                 "Segurança Auth JWT",
                 "SQLAlchemy ORM (SQLite)",
                 "Document & Policy Services",
                 "API OpenAPI (Swagger)"
             ], ACCENT_BLUE)

    add_card(slide4, Inches(6.8), Inches(1.8), Inches(2.7), Inches(5.0), 
             "Motor Agêntico & RAG", 
             [
                 "LangGraph Orchestrator",
                 "Pydantic Agent v2",
                 "PyMuPDF OCR Engine",
                 "ChromaDB Vector Store",
                 "Gemini 2.5 LLM via OpenRouter"
             ], ACCENT_BLUE)

    add_card(slide4, Inches(9.8), Inches(1.8), Inches(2.7), Inches(5.0), 
             "Ferramentas de Suporte", 
             [
                 "Antigravity (Pair Programming Agent)",
                 "Google Stitch (UI/UX Prototyping)",
                 "NotebookLM (Síntese & Roteirização)",
                 "CapCut (Edição & Legendagem)",
                 "Claude AI & ChatGPT (Prompts/Dados)"
             ], ACCENT_BLUE)

    # -------------------------------------------------------------
    # SLIDE 5: ECOSSISTEMA DE AGENTES
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide5)
    add_header(slide5, "Motor Agêntico & IA", "Ecossistema de 5 Agentes de IA Especializados")

    col_w = Inches(2.2)
    spacing = Inches(0.25)
    left_start = Inches(0.8)

    add_card(slide5, left_start, Inches(1.8), col_w, Inches(5.0), 
             "1. DocumentParser", 
             [
                 "Ingestão do PDF.",
                 "Validação SHA-256.",
                 "OCR PyMuPDF.",
                 "Extração de layout por página."
             ], ACCENT_CYAN)

    add_card(slide5, left_start + (col_w + spacing)*1, Inches(1.8), col_w, Inches(5.0), 
             "2. PolicyExtraction", 
             [
                 "Parsing semântico.",
                 "Schemas Pydantic estritos.",
                 "Extração de LMR, Prêmio, Franquias Side A/B/C.",
                 "Classificação de Cláusulas."
             ], ACCENT_CYAN)

    add_card(slide5, left_start + (col_w + spacing)*2, Inches(1.8), col_w, Inches(5.0), 
             "3. EvidenceGrounded", 
             [
                 "Vinculação de provas.",
                 "Mapeamento do texto exato no PDF.",
                 "Numeração de página auditável.",
                 "Garantia anti-alucinação."
             ], ACCENT_CYAN)

    add_card(slide5, left_start + (col_w + spacing)*3, Inches(1.8), col_w, Inches(5.0), 
             "4. PolicyComparison", 
             [
                 "Orquestrado em LangGraph.",
                 "Alinhamento semântico Apólice A vs B.",
                 "Cálculo de Deltas.",
                 "Classificação de Severidade de Risco."
             ], ACCENT_CYAN)

    add_card(slide5, left_start + (col_w + spacing)*4, Inches(1.8), col_w, Inches(5.0), 
             "5. ReportGeneration", 
             [
                 "Consolidação de insights.",
                 "Renderização de tabelas e gráficos.",
                 "Exportação HTML/PDF.",
                 "Geração de relatórios executivos."
             ], ACCENT_CYAN)

    # -------------------------------------------------------------
    # SLIDE 6: UI/UX
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide6)
    add_header(slide6, "Experiência de Utilizador (UI/UX)", "Interface B2B Auditável com Feedback Transparente")

    add_card(slide6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.35), 
             "Modal de Feedback Visual (AI Overlay)", 
             [
                 "Overlay animado com spinner e barra de progresso durante o processamento.",
                 "Comunicação clara dos passos técnicos da IA (OCR, RAG, Extração, Comparação).",
                 "Garante transparência operacional e elimina incertezas no upload de arquivos grandes."
             ], ACCENT_BLUE)

    add_card(slide6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.35), 
             "Visualizador Split-Screen Auditável", 
             [
                 "Exibição simultânea do PDF original e dos metadados extraídos.",
                 "Navegação automática para a página exata da evidência textual.",
                 "Abas dedicadas para Coberturas, Exclusões, Limites, Franquias e Cláusulas."
             ], ACCENT_BLUE)

    add_card(slide6, Inches(0.8), Inches(4.45), Inches(5.6), Inches(2.35), 
             "Matriz Comparativa & Gráficos Radar", 
             [
                 "Visualização lado a lado da Apólice Base vs Comparada.",
                 "Gráficos de Radar de cobertura e gráficos de distribuição de severidade de risco.",
                 "Filtros interativos por severidade (Alto, Médio, Baixo Risco)."
             ], ACCENT_BLUE)

    add_card(slide6, Inches(6.8), Inches(4.45), Inches(5.7), Inches(2.35), 
             "Histórico de Análises & Exportação", 
             [
                 "Rastreamento completo de análises e comparações com ID, Data e Hora.",
                 "Filtros por seguradora, limite e palavras-chave de cláusulas.",
                 "Exportação imediata de relatórios executivos em HTML/PDF."
             ], ACCENT_BLUE)

    # -------------------------------------------------------------
    # SLIDE 7: IMPACTO & RESULTADOS
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide7)
    add_header(slide7, "Resultados & Ganho de Eficiência", "Impacto de Negócio Medível para Seguradoras e Corretoras")

    add_card(slide7, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "> 95% Redução de Tempo", 
             [
                 "De 4 horas de análise manual para menos de 30 segundos por apólice.",
                 "Aceleração do processo de cotação e emissão de seguros.",
                 "Resposta rápida a clientes corporativos em disputas de coberturas."
             ], ACCENT_EMERALD)

    add_card(slide7, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "100% Auditabilidade & Compliance", 
             [
                 "Eliminação do risco de alucinações através de Grounded Proof.",
                 "Conformidade total com auditorias de underwriting e sinistros.",
                 "Garantia de evidência documental com numeração de página exata."
             ], ACCENT_EMERALD)

    add_card(slide7, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "10x Produtividade Operacional", 
             [
                 "Capacidade de analisar 10x mais apólices com a mesma equipa.",
                 "Redução drástica de falhas humanas na identificação de exclusões.",
                 "Padronização técnica nas análises de risco D&O."
             ], ACCENT_EMERALD)

    # -------------------------------------------------------------
    # SLIDE 8: ROTEIRO FUTURO
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide8)
    add_header(slide8, "Plano de Evolução Futura", "Roteiro Tecnológico & Expansão do Produto")

    add_card(slide8, Inches(0.8), Inches(1.8), Inches(2.7), Inches(5.0), 
             "Fase 1 (Atual)", 
             [
                 "MVP D&O Concluído.",
                 "5 Agentes IA + RAG.",
                 "Split-Screen UI.",
                 "SQLite + FastAPI.",
                 "Grounded Proof 100%."
             ], ACCENT_GOLD)

    add_card(slide8, Inches(3.8), Inches(1.8), Inches(2.7), Inches(5.0), 
             "Fase 2 (Curto Prazo)", 
             [
                 "Ontologia Multi-Ramo.",
                 "Suporte a E&O, Cyber Risks e Garantia.",
                 "Suporte a OCR de fotocópias com OpenCV."
             ], ACCENT_GOLD)

    add_card(slide8, Inches(6.8), Inches(1.8), Inches(2.7), Inches(5.0), 
             "Fase 3 (Médio Prazo)", 
             [
                 "Integração com Core Systems.",
                 "Conetores para Guidewire, SAP e Salesforce.",
                 "WebSockets em tempo real para streaming de análise."
             ], ACCENT_GOLD)

    add_card(slide8, Inches(9.8), Inches(1.8), Inches(2.7), Inches(5.0), 
             "Fase 4 (Longo Prazo)", 
             [
                 "Suporte a SLMs On-Premise.",
                 "Execução de Llama 3 8B / Mistral em servidores locais.",
                 "Garantia de sigilo absoluto e compliance bancário/segurador."
             ], ACCENT_GOLD)

    # -------------------------------------------------------------
    # SLIDE 9: ENCERRAMENTO
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide9)

    card9 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.1))
    card9.fill.solid()
    card9.fill.fore_color.rgb = CARD_BG
    card9.line.color.rgb = CARD_BORDER
    card9.line.width = Pt(1.5)

    tb9 = slide9.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.333), Inches(4.0))
    tf9 = tb9.text_frame
    tf9.word_wrap = True

    p = tf9.paragraphs[0]
    p.text = "PROJETO FINAL — INSURMINDS (I2A2)  |  GRUPO IA JOURNEY"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(14)

    p2 = tf9.add_paragraph()
    p2.text = "Obrigado!"
    p2.font.size = Pt(44)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_LIGHT
    p2.space_after = Pt(10)

    p3 = tf9.add_paragraph()
    p3.text = "D&O Intelligence — Transformando a Análise e Comparação de Apólices com Inteligência Artificial Agêntica"
    p3.font.size = Pt(18)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_after = Pt(30)

    p4 = tf9.add_paragraph()
    p4.text = "Grupo: IA Journey  |  Autor: Arsénio António Monjane  |  monjearse@hotmail.com  |  Outubro 2026"
    p4.font.size = Pt(14)
    p4.font.bold = True
    p4.font.color.rgb = ACCENT_CYAN

    # Save presentation
    output_path = r"d:\Projectos\I2A2\Labs\DO Insight — Intelligent Policy Analysis\InsurMinds_Projeto_Final.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
